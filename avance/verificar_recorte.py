# -*- coding: utf-8 -*-
"""Red de seguridad del recorte (A-04 del dictamen 12).

Compara dos versiones del .docx y avisa si el recorte se llevó algo que no debía:
  - una cifra (decimal, porcentaje, estadístico) que estaba y ya no aparece en ningún lugar;
  - una cita «Autor (año)» / «(Autor, año)» que desaparece del todo (quedaría huérfana);
  - una remisión a Tabla/Figura/Anexo que deja de existir como remisión;
  - una remisión «§x.y» que apunta a un apartado que no existe en la versión nueva.
Además informa las palabras por capítulo, antes y después.

No juzga si el texto quedó bien: sólo garantiza que no se perdió información verificable.

Uso: python verificar_recorte.py "<antes.docx>" "<después.docx>"
"""
import re
import sys
from collections import Counter

import docx

EXCLUIR = ('table of figures', 'TOC Heading')


def cuerpo(ruta):
    doc = docx.Document(ruta)
    textos, caps, cap = [], Counter(), 'preliminar'
    for p in doc.paragraphs:
        if p.style.name in EXCLUIR or p.style.name.startswith('toc'):
            continue
        if p.style.name == 'Heading 1':
            cap = p.text.strip()[:40]
        textos.append(p.text)
        caps[cap] += len(p.text.split())
    for t in doc.tables:
        for r in t.rows:
            for c in r.cells:
                textos.append(c.text)
    titulos = [p.text.strip() for p in doc.paragraphs if p.style.name.startswith('Heading')]
    return '\n'.join(textos), caps, titulos


CIFRA = re.compile(r'(?<![\w/])(?:\d+,\d+\s*%?|\d+\s*%|[tdnpκα]\s*(?:\(\d+\))?\s*[=<>]\s*-?\d+(?:,\d+)?)')
CITA = re.compile(r"([A-ZÁÉÍÓÚÑ][\w'´\-]+(?: (?:y|&) [A-ZÁÉÍÓÚÑ][\w'´\-]+| et al\.)?),? \(?((?:19|20)\d\d|s\.f\.)")
REMI = re.compile(r'\b(Tabla \d+|Figura \d+|Anexo [A-E](?:\.\d+)*)')
SECC = re.compile(r'§(\d+(?:\.\d+)*(?:\.[a-z])?)')


def main(a, b):
    ta, ca, _ = cuerpo(a)
    tb, cb, tit_b = cuerpo(b)
    fallos = 0

    def perdidos(pat, nombre, norm=lambda m: m):
        nonlocal fallos
        sa = Counter(norm(m) for m in pat.findall(ta))
        sb = Counter(norm(m) for m in pat.findall(tb))
        idos = sorted(k for k in sa if k not in sb)
        print(f'\n── {nombre}: {len(sa)} distintos antes, {len(sb)} después, '
              f'{len(idos)} desaparecen del todo')
        for k in idos:
            print('   *', k)
        fallos += len(idos)

    perdidos(CIFRA, 'cifras', lambda m: re.sub(r'\s+', '', m))
    perdidos(CITA, 'citas', lambda m: f'{m[0]} {m[1]}')
    perdidos(REMI, 'remisiones a tabla / figura / anexo')

    numeros = {re.match(r'[\d.]+[a-z]?', t).group(0).rstrip('.') for t in tit_b
               if re.match(r'\d', t)}
    anexos = {re.match(r'([A-E](?:\.\d+)*)', t).group(1) for t in tit_b if re.match(r'[A-E]\.\d', t)}
    rotas = sorted({s for s in SECC.findall(tb)
                    if s not in numeros and not any(n.startswith(s + '.') for n in numeros)
                    and not re.match(r'\d+\.\d+\.[a-z]$', s) and not s.endswith('.')})
    print(f'\n── remisiones «§» a apartados inexistentes: {len(rotas)}')
    for s in rotas:
        print('   *', s)

    print('\n── palabras por capítulo (antes → después)')
    tot_a = tot_b = 0
    for k in list(ca) + [k for k in cb if k not in ca]:
        print(f'   {k[:38]:<38} {ca.get(k, 0):>7} → {cb.get(k, 0):>7}  ({cb.get(k, 0) - ca.get(k, 0):+d})')
        tot_a += ca.get(k, 0)
        tot_b += cb.get(k, 0)
    print(f'   {"TOTAL":<38} {tot_a:>7} → {tot_b:>7}  ({tot_b - tot_a:+d})')
    print('\nsin pérdidas' if fallos == 0 else f'\n*** {fallos} elementos desaparecieron: revisar ***')
    return 1 if fallos else 0


if __name__ == '__main__':
    sys.exit(main(sys.argv[1], sys.argv[2]))

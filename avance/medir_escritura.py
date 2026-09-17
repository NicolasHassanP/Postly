# -*- coding: utf-8 -*-
"""
Métricas de escritura del documento (N-22 del dictamen).

Reproduce la medición del dictamen: segmenta por párrafo y luego por final de oración,
excluyendo los índices automáticos, los rótulos de tabla y figura y las filas de tabla.
Separa cuerpo (Resumen → Cap. 6) de anexos, porque el dictamen los reporta por separado.

Uso: python medir_escritura.py "<archivo.docx>" [--listar N]
"""
import re
import sys

import docx

EXCLUIR = ('table of figures', 'TOC Heading')
ROTULO = re.compile(r'^(Tabla|Figura|Nota\.)\s')
# fin de oración: puntuación fuerte seguida de mayúscula, apertura de cita, numeral de
# enumeración —«(1)», «1.»— o un apellido de partícula minúscula («vom Brocke», «van …»)
FIN = re.compile(r'(?<=[.!?…])\s+(?=[«"¿¡(\d]|[A-ZÁÉÍÓÚÑ]|(?:vom|van|de|della)\s+[A-Z])')


def oraciones(texto):
    return [o.strip() for o in FIN.split(texto) if len(o.split()) > 2]


def bloques(doc):
    """(cuerpo, anexos) como listas de párrafos legibles."""
    idx_ref = next(i for i, p in enumerate(doc.paragraphs)
                   if p.style.name == 'Heading 1' and 'REFERENCIAS' in p.text.upper())
    idx_anx = next(i for i, p in enumerate(doc.paragraphs)
                   if p.style.name == 'Heading 1' and 'ANEXOS' in p.text.upper())
    idx_res = next(i for i, p in enumerate(doc.paragraphs)
                   if p.style.name == 'Heading 1' and p.text.strip().upper() == 'RESUMEN')

    def limpiar(rango):
        out = []
        for i in rango:
            p = doc.paragraphs[i]
            if p.style.name in EXCLUIR or p.style.name.startswith('toc'):
                continue
            if p.style.name.startswith('Heading'):
                continue
            t = p.text.strip()
            if not t or ROTULO.match(t):
                continue
            out.append(t)
        return out

    return limpiar(range(idx_res, idx_ref)), limpiar(range(idx_anx, len(doc.paragraphs)))


def informe(nombre, parrafos, listar=0):
    ors = [o for p in parrafos for o in oraciones(p)]
    largos = [len(o.split()) for o in ors]
    if not largos:
        return []
    n50 = [o for o in ors if len(o.split()) > 50]
    n40 = [o for o in ors if len(o.split()) > 40]
    print(f'\n── {nombre}')
    print(f'   palabras                 {sum(largos)}')
    print(f'   oraciones                {len(ors)}')
    print(f'   media de palabras/oración{sum(largos) / len(largos):8.1f}')
    print(f'   oraciones > 40 palabras  {len(n40)} ({100 * len(n40) / len(ors):.1f} %)')
    print(f'   oraciones > 50 palabras  {len(n50)} ({100 * len(n50) / len(ors):.1f} %)')
    print(f'   oración más larga        {max(largos)} palabras')
    peores = sorted(ors, key=lambda o: -len(o.split()))[:listar]
    for o in peores:
        print(f'\n   [{len(o.split())}] {o}')
    return n50


if __name__ == '__main__':
    doc = docx.Document(sys.argv[1])
    listar = int(sys.argv[sys.argv.index('--listar') + 1]) if '--listar' in sys.argv else 0
    cuerpo, anexos = bloques(doc)
    informe('CUERPO (Resumen → Cap. 6)', cuerpo, listar)
    informe('ANEXOS', anexos, listar)
    print('\n   objetivo del plan de remediación: 25-28 palabras por oración en el cuerpo')
